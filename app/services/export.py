from __future__ import annotations

from datetime import date
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from ..models import Report
from .reports import month_label


def build_workbook(first: date, reports: list[Report]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Отчёты"

    headers = ["Дата", "Сотрудник", "Задача", "Часы", "Комментарий"]
    sheet.append(headers)
    for column in range(1, len(headers) + 1):
        sheet.cell(row=1, column=column).font = Font(bold=True)

    total = 0.0
    for report in sorted(reports, key=lambda r: (r.report_date, r.user.full_name)):
        for item in report.items:
            hours = float(item.hours)
            total += hours
            sheet.append(
                [
                    report.report_date,
                    report.user.full_name,
                    item.task_title,
                    hours,
                    item.comment or "",
                ]
            )

    last_row = sheet.max_row + 1
    sheet.cell(row=last_row, column=3, value="Итого").font = Font(bold=True)
    sheet.cell(row=last_row, column=4, value=round(total, 2)).font = Font(bold=True)

    widths = [12, 28, 60, 8, 40]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for row in sheet.iter_rows(min_row=2, min_col=1, max_col=1):
        for cell in row:
            cell.number_format = "DD.MM.YYYY"
    for row in sheet.iter_rows(min_row=2, min_col=3, max_col=3):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    sheet.freeze_panes = "A2"

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def export_filename(first: date) -> str:
    return f"reports-{first:%Y-%m}.xlsx"
